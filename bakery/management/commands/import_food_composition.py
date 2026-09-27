"""日本食品標準成分表（八訂）の本表（Excel）を取り込む。何度実行しても重複しない（食品番号で上書き）。

使い方:
    python manage.py import_food_composition            # data/mext_food_composition_8th.xlsx を読む
    python manage.py import_food_composition --file 別のファイル.xlsx

出典：日本食品標準成分表（八訂）増補2023年（文部科学省）
    https://www.mext.go.jp/a_menu/syokuhinseibun/mext_00001.html
列は見出しの「成分識別子」（ENERC_KCAL など）で探すので、列の位置が変わっても読める。
"""

from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from bakery.models import FoodComposition

DEFAULT_FILE = Path(settings.BASE_DIR) / "data" / "mext_food_composition_8th.xlsx"
SHEET = "表全体"
COLUMNS = {  # 成分識別子 → モデルの項目
    "ENERC_KCAL": "energy_kcal",
    "PROT-": "protein_g",
    "FAT-": "fat_g",
    "CHOCDF-": "carbohydrate_g",
    "NACL_EQ": "salt_g",
}


def parse_value(value):
    """成分表の値 → Decimal。Tr（微量）は0、括弧付きの推定値は括弧内、「-」「*」や空欄は None。"""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().strip("()（）")
    if text in ("Tr", "(Tr)"):
        return Decimal("0")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


class Command(BaseCommand):
    help = "日本食品標準成分表（八訂）の本表を取り込む"

    def add_arguments(self, parser):
        parser.add_argument("--file", default=str(DEFAULT_FILE), help="成分表の Excel ファイル")

    @transaction.atomic
    def handle(self, *args, **options):
        import openpyxl

        path = Path(options["file"])
        if not path.exists():
            raise CommandError(f"ファイルがありません：{path}")
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        # シート名「1穀類」「2いも及びでん粉類」… から食品群の名前を作る（01 → 穀類）
        groups = {}
        for name in workbook.sheetnames:
            digits = "".join(ch for ch in name if ch.isdigit())
            if digits:
                groups[digits.zfill(2)] = name[len(digits):]
        rows = workbook[SHEET].iter_rows(values_only=True)

        index = None
        for row in rows:
            if "ENERC_KCAL" in row:  # 成分識別子の行
                index = {field: row.index(code) for code, field in COLUMNS.items()}
                break
        if index is None:
            raise CommandError("成分識別子（ENERC_KCAL など）の行が見つかりません。ファイルを確認してください")

        count = 0
        for row in rows:
            number = str(row[1] or "").strip()
            if not number.isdigit():
                continue
            values = {field: parse_value(row[i]) for field, i in index.items()}
            FoodComposition.objects.update_or_create(
                food_number=number,
                defaults={"group": groups.get(str(row[0] or "").strip().zfill(2), str(row[0] or "")),
                          "name": " ".join(str(row[3]).split()), **values},
            )
            count += 1
        self.stdout.write(self.style.SUCCESS(f"成分表を取り込みました：{count} 食品"))
