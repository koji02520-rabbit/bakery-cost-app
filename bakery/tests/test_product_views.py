"""商品画面（S30〜S33）のテスト。材料・レシピの準備は test_recipe_views の RecipeTestBase を使う。"""

from decimal import Decimal

from django.urls import reverse

from bakery.models import AuditLog, CostSnapshot, Product, Recipe, RecipeAdoption, SnapshotReason, Status
from bakery.tests.test_recipe_views import RecipeTestBase


class ProductTestBase(RecipeTestBase):
    def make_final(self):
        """生地（強力粉1000g＋バター100g）60g を使うバターロール。原価 22.26…円 → 切り上げて 22.3円。"""
        dough = self.create("パン生地")
        self.add(dough, self.flour, 1000, "1")
        self.add(dough, self.butter, 100, "2")
        bread = self.create("バターロール", kind="final", recipe_type="")
        self.add(bread, dough, 60)
        return dough, bread

    def register(self, name="バターロール", price="200", recipe=None):
        self.client.force_login(self.admin)
        return self.client.post(reverse("product_new"), {
            "name": name, "price_excluding_tax": price, "adopted_recipe": recipe.pk if recipe else "",
        })


class RegisterTests(ProductTestBase):
    def test_register_with_recipe(self):
        dough, bread = self.make_final()
        response = self.register(recipe=bread)
        product = Product.objects.get()
        self.assertRedirects(response, reverse("product_detail", args=[product.pk]))
        self.assertEqual((product.code, product.adopted_recipe, product.created_by), ("ITEM-001", bread, self.admin))
        adoption = RecipeAdoption.objects.get(product=product)
        self.assertIsNone(adoption.old_recipe)
        self.assertEqual(adoption.new_cost, Decimal("22.3"))
        self.assertEqual(CostSnapshot.objects.get(product=product).reason, SnapshotReason.ADOPTION)
        self.assertTrue(AuditLog.objects.filter(action="商品登録", target_id=product.pk).exists())

        detail = self.client.get(reverse("product_detail", args=[product.pk]))
        for text in ["200 円", "216 円", "22.3 円", "11.2%", "小麦", "乳", "パン生地 → 強力粉"]:
            self.assertContains(detail, text)

    def test_register_without_recipe(self):
        self.register(name="試作品")
        product = Product.objects.get()
        self.assertIsNone(product.adopted_recipe)
        self.assertContains(self.client.get(reverse("product_list")), "未設定")

    def test_price_input(self):
        for price, message in [("", "税抜売価を入力してください"), ("12.5", "税抜売価は整数（円）で入力してください"),
                               ("-1", "税抜売価は0円以上で入力してください")]:
            with self.subTest(price=price):
                self.assertContains(self.register(price=price), message)
        self.register(price="２８０円")
        self.assertEqual(Product.objects.get().price_excluding_tax, 280)

    def test_tax_is_rounded_down(self):
        self.register(price="155")  # 155 × 1.08 = 167.4 → 167
        self.assertContains(self.client.get(reverse("product_list")), "167 円")

    def test_only_final_recipes_with_items_can_be_adopted(self):
        dough, bread = self.make_final()
        empty = self.create("空のパン", kind="final", recipe_type="")
        for recipe in (dough, empty):
            with self.subTest(recipe=recipe.name):
                self.assertContains(self.register(recipe=recipe), "採用レシピを選び直してください")
        self.assertFalse(Product.objects.exists())


class SwitchTests(ProductTestBase):
    def setUp(self):
        super().setUp()
        self.dough, self.bread = self.make_final()
        self.register(recipe=self.bread)
        self.product = Product.objects.get()
        # 新しいバージョン：生地を 80g に増やす → 29.68…円 → 切り上げて 29.7円
        self.client.force_login(self.staff)
        url = self.client.post(reverse("recipe_new_version", args=[self.bread.pk])).url
        self.v2 = Recipe.objects.get(pk=url.split("/")[2])
        row = self.v2.items.get()
        self.client.post(reverse("recipe_item_update", args=[row.pk]), {"quantity_g": "80", "step_label": ""})
        self.client.force_login(self.admin)

    def test_detail_suggests_newer_version(self):
        detail = self.client.get(reverse("product_detail", args=[self.product.pk]))
        self.assertContains(detail, "採用レシピに新しいバージョンがあります")

    def test_preview_then_switch(self):
        url = reverse("product_switch", args=[self.product.pk])
        preview = self.client.get(url, {"recipe": self.v2.pk})
        for text in ["22.3 円", "29.7 円", "+7.4 円", "11.2%", "14.9%"]:
            self.assertContains(preview, text)
        self.product.refresh_from_db()
        self.assertEqual(self.product.adopted_recipe, self.bread)  # 確認画面ではまだ変えない

        response = self.client.post(url, {"recipe": self.v2.pk})
        self.assertRedirects(response, reverse("product_detail", args=[self.product.pk]))
        self.product.refresh_from_db()
        self.assertEqual(self.product.adopted_recipe, self.v2)
        adoption = self.product.adoptions.first()
        self.assertEqual((adoption.old_recipe, adoption.new_recipe), (self.bread, self.v2))
        self.assertEqual(adoption.old_cost, Decimal("22.3"))
        self.assertEqual(self.product.cost_snapshots.count(), 2)
        log = AuditLog.objects.get(action="採用レシピ切替")
        self.assertEqual(log.after["adopted_recipe"]["id"], self.v2.pk)
        # 旧バージョンはどこにも採用されなくなる → 下書き扱いで編集できる
        self.assertFalse(self.bread.is_in_use())

    def test_cannot_switch_to_same_or_stopped(self):
        url = reverse("product_switch", args=[self.product.pk])
        self.assertContains(self.client.get(url, {"recipe": self.bread.pk}), "今の採用レシピと同じレシピ")
        self.client.post(reverse("recipe_set_status", args=[self.v2.pk]), {"status": "stopped"})
        self.assertContains(self.client.post(url, {"recipe": self.v2.pk}), "使用停止中のため採用できません")
        self.product.refresh_from_db()
        self.assertEqual(self.product.adopted_recipe, self.bread)


class ListAndPermissionTests(ProductTestBase):
    def test_sort_by_cost_rate(self):
        dough, bread = self.make_final()
        self.register(name="高い商品", price="100", recipe=bread)  # 22.3%
        self.register(name="安い商品", price="400", recipe=bread)  # 5.6%
        response = self.client.get(reverse("product_list"), {"sort": "rate"})
        body = response.content.decode()
        self.assertLess(body.index("高い商品"), body.index("安い商品"))
        response = self.client.get(reverse("product_list"), {"sort": "code"})
        self.assertContains(response, "22.3%")

    def test_edit_name_and_price(self):
        dough, bread = self.make_final()
        self.register(recipe=bread)
        product = Product.objects.get()
        self.client.post(reverse("product_edit", args=[product.pk]), {"name": "バターロール（大）", "price_excluding_tax": "250"})
        product.refresh_from_db()
        self.assertEqual((product.name, product.price_excluding_tax, product.adopted_recipe), ("バターロール（大）", 250, bread))
        log = AuditLog.objects.get(action="商品修正")
        self.assertEqual((log.before["price_excluding_tax"], log.after["price_excluding_tax"]), (200, 250))

    def test_general_user_can_view_only(self):
        dough, bread = self.make_final()
        self.register(recipe=bread)
        product = Product.objects.get()
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(reverse("product_list")).status_code, 200)
        detail = self.client.get(reverse("product_detail", args=[product.pk]))
        self.assertEqual(detail.status_code, 200)
        self.assertNotContains(detail, "採用レシピを切り替える")
        for name in ["product_edit", "product_switch"]:
            self.assertEqual(self.client.get(reverse(name, args=[product.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse("product_new")).status_code, 403)
        self.assertEqual(self.client.post(reverse("product_set_status", args=[product.pk]), {"status": "stopped"}).status_code, 403)

    def test_stop_product_and_pages_open(self):
        dough, bread = self.make_final()
        self.register(recipe=bread)
        product = Product.objects.get()
        self.client.post(reverse("product_set_status", args=[product.pk]), {"status": "stopped"}, follow=True)
        product.refresh_from_db()
        self.assertEqual(product.status, Status.STOPPED)
        self.assertNotContains(self.client.get(reverse("product_list")), "ITEM-001")
        self.assertContains(self.client.get(reverse("product_list"), {"status": "stopped"}), "ITEM-001")
        for name, args in [("product_new", []), ("product_detail", [product.pk]), ("product_edit", [product.pk]),
                           ("product_switch", [product.pk]), ("recipe_detail", [bread.pk]),
                           ("ingredient_detail", [self.flour.pk])]:
            with self.subTest(page=name):
                self.assertEqual(self.client.get(reverse(name, args=args)).status_code, 200)
        self.assertContains(self.client.get(reverse("home"), {"q": "バターロール"}), "ITEM-001")


class ExportTests(ProductTestBase):
    """商品原価表のエクスポート（第68項）。"""

    def setUp(self):
        super().setUp()
        dough, bread = self.make_final()
        self.register(name="バターロール", price="200", recipe=bread)  # 原価 22.3円・税込 216円
        self.register(name="=試作品", price="155")  # レシピ未設定・税込 167円
        self.client.force_login(self.staff)  # 閲覧できる人なら誰でも出力できる

    def export(self, fmt, **params):
        return self.client.get(reverse("product_export", args=[fmt]), params)

    def test_csv(self):
        response = self.export("csv")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn(".csv", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))  # Excel で文字化けしないよう BOM 付き
        self.assertEqual(response.content.decode("utf-8-sig").splitlines(), [
            "商品ID,商品名,原価（円）,売価（税抜・円）,税込売価（円）",
            "ITEM-001,バターロール,22.3,200,216",
            "ITEM-002,'=試作品,,155,167",  # 原価を計算できない商品は空欄、数式に見える名前は文字として出す
        ])

    def test_xlsx(self):
        from io import BytesIO

        from openpyxl import load_workbook

        response = self.export("xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertIn(".xlsx", response["Content-Disposition"])
        sheet = load_workbook(BytesIO(response.content)).active
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        self.assertEqual(rows, [
            ["商品ID", "商品名", "原価（円）", "売価（税抜・円）", "税込売価（円）"],
            ["ITEM-001", "バターロール", 22.3, 200, 216],
            ["ITEM-002", "=試作品", None, 155, 167],
        ])
        self.assertEqual(sheet["B3"].data_type, "s")  # 数式にしない
        self.assertEqual(sheet["C2"].number_format, "#,##0.0")

    def test_follows_list_filters_and_sort(self):
        lines = self.export("csv", q="バター").content.decode("utf-8-sig").splitlines()
        self.assertEqual([line.split(",")[0] for line in lines[1:]], ["ITEM-001"])
        lines = self.export("csv", sort="cost").content.decode("utf-8-sig").splitlines()
        self.assertEqual([line.split(",")[0] for line in lines[1:]], ["ITEM-001", "ITEM-002"])

        Product.objects.filter(code="ITEM-001").update(status=Status.STOPPED)
        self.assertNotIn("ITEM-001", self.export("csv").content.decode("utf-8-sig"))
        self.assertIn("ITEM-001", self.export("csv", status="all").content.decode("utf-8-sig"))

    def test_list_has_links_with_current_filters(self):
        response = self.client.get(reverse("product_list"), {"q": "バター", "sort": "rate"})
        self.assertContains(response, reverse("product_export", args=["csv"]) + "?q=")
        self.assertContains(response, "sort=rate")

    def test_unknown_format_and_login(self):
        self.assertEqual(self.export("pdf").status_code, 404)
        self.client.logout()
        self.assertEqual(self.export("csv").status_code, 302)
