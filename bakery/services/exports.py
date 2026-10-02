"""商品原価表のエクスポート（CSV・Excel）（spec.md 第68項）。

項目は 商品ID・商品名・原価・税抜売価・税込売価 の5つ。原価は画面と同じく小数第1位まで。
原価を計算できない商品（採用レシピなし等）は空欄にする（0 と見せない）。
"""

from __future__ import annotations

import csv
import io
from decimal import ROUND_HALF_UP, Decimal

HEADERS = ["商品ID", "商品名", "原価（円）", "売価（税抜・円）", "税込売価（円）"]

# 表計算ソフトが数式として解釈する先頭文字（CSV インジェクション対策）
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def cost_rows(product_costs):
    """ProductCost の並び → [商品ID, 商品名, 原価(Decimal|None), 税抜売価, 税込売価] の並び。"""
    rows = []
    for pc in product_costs:
        cost = None if pc.cost is None else pc.cost.quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
        rows.append([pc.product.code, pc.product.name, cost, pc.product.price_excluding_tax, pc.price_including_tax])
    return rows


def _safe_text(value):
    return "'" + value if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES) else value


def to_csv(rows) -> bytes:
    """Excel でそのまま開けるよう、BOM 付き UTF-8・CRLF にする。"""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(HEADERS)
    for row in rows:
        writer.writerow(["" if v is None else _safe_text(v) for v in row])
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(rows) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "商品原価"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for code, name, cost, price, price_with_tax in rows:
        sheet.append([code, name, cost, price, price_with_tax])
        # 商品名が「=」で始まっても数式にせず文字として入れる
        sheet.cell(row=sheet.max_row, column=2).data_type = "s"
    for row in sheet.iter_rows(min_row=2):
        row[2].number_format = "#,##0.0"
        row[3].number_format = row[4].number_format = "#,##0"
    for column, width in zip("ABCDE", [12, 30, 12, 16, 14]):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
